from __future__ import annotations

from datetime import date
from decimal import Decimal

from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from analyzer.models.movement import Movement


class WorksheetWriter:
    """
    Excel çalışma sayfalarını doldurur.
    """

    def __init__(self) -> None:

        self._title_font = Font(
            bold=True,
            size=16,
        )

        self._header_font = Font(
            bold=True,
        )

    def _write_title(
        self,
        worksheet: Worksheet,
        title: str,
    ) -> None:
        """
        Sayfanın başlığını yazar.
        """

        cell = worksheet["A1"]

        cell.value = title
        cell.font = self._title_font

    def _write_headers(
        self,
        worksheet: Worksheet,
        headers: list[str],
        row: int = 3,
    ) -> None:
        """
        Sütun başlıklarını yazar.
        """

        for column, header in enumerate(
            headers,
            start=1,
        ):
            cell = worksheet.cell(
                row=row,
                column=column,
            )

            cell.value = header
            cell.font = self._header_font

    def _movement_row(
        self,
        movement: Movement,
    ) -> list[object]:
        """
        Movement nesnesini Excel satırına dönüştürür.
        """

        return [
            movement.tif_no,
            movement.movement_date,
            movement.warehouse,
            movement.budget_type,
            movement.stock_code,
            movement.stock_name,
            movement.supplier,
            movement.quantity,
            movement.amount,
        ]

    def _auto_fit_columns(
        self,
        worksheet: Worksheet,
    ) -> None:
        """
        Sütun genişliklerini otomatik ayarlar.
        """

        for column_cells in worksheet.columns:

            values = [
                str(cell.value)
                for cell in column_cells
                if cell.value is not None
            ]

            if not values:
                continue

            length = max(
                len(value)
                for value in values
            )

            worksheet.column_dimensions[
                column_cells[0].column_letter
            ].width = min(
                length + 2,
                50,
            )

    def _format_date(
        self,
        value: date | None,
    ) -> str:

        if value is None:
            return ""

        return value.strftime(
            "%d.%m.%Y"
        )

    def _format_decimal(
        self,
        value: Decimal | None,
    ) -> str:

        if value is None:
            return ""

        return f"{value:.2f}"
    

    def write_summary(
        self,
        worksheet: Worksheet,
        result,
    ) -> None:
        """
        Özet sayfasını oluşturur.
        """

        self._write_title(
            worksheet,
            "MKYS - TDMS Uzlaştırma Raporu",
        )

        headers = [
            "Kategori",
            "Adet",
        ]

        self._write_headers(
            worksheet,
            headers,
        )

        rows = [
            (
                "Eşleşen Giriş Hareketleri",
                len(result.matched),
            ),
            (
                "TDMS'de Eksik",
                len(result.missing_in_tdms),
            ),
            (
                "MKYS'de Eksik",
                len(result.missing_in_mkys),
            ),
            (
                "Tutar Farkları",
                len(result.amount_differences),
            ),
            (
                "Tüketim Farkları",
                len(result.consumption_differences),
            ),
            (
                "Açılış Eşleşmeleri",
                len(result.opening_matched),
            ),
            (
                "TDMS'de Eksik Açılış",
                len(result.opening_missing_in_tdms),
            ),
            (
                "MKYS'de Eksik Açılış",
                len(result.opening_missing_in_mkys),
            ),
        ]

        row_number = 4

        for title, value in rows:

            worksheet.cell(
                row=row_number,
                column=1,
            ).value = title

            worksheet.cell(
                row=row_number,
                column=2,
            ).value = value

            row_number += 1

        self._auto_fit_columns(
            worksheet,
        )

    def write_movements(
        self,
        worksheet: Worksheet,
        title: str,
        movements: list[Movement],
    ) -> None:
        """
        Hareket listesini çalışma sayfasına yazar.
        """

        self._write_title(
            worksheet,
            title,
        )

        headers = [
            "TİF No",
            "Tarih",
            "Depo",
            "Bütçe",
            "Taşınır Kodu",
            "Malzeme",
            "Tedarikçi",
            "Miktar",
            "Tutar",
        ]

        self._write_headers(
            worksheet,
            headers,
        )

        row_number = 4

        for movement in movements:

            values = self._movement_row(
                movement,
            )

            for column, value in enumerate(
                values,
                start=1,
            ):

                cell = worksheet.cell(
                    row=row_number,
                    column=column,
                )

                if isinstance(
                    value,
                    date,
                ):
                    cell.value = self._format_date(
                        value,
                    )

                elif isinstance(
                    value,
                    Decimal,
                ):
                    cell.value = float(value)

                else:
                    cell.value = value

            row_number += 1

        self._auto_fit_columns(
            worksheet,
        )