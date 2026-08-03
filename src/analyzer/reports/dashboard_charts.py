from __future__ import annotations

from openpyxl.chart import (
    BarChart,
    PieChart,
    Reference,
)
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.worksheet.worksheet import Worksheet

from analyzer.reports.dashboard_summary import DashboardSummary

CHART_HEIGHT = 7
CHART_WIDTH = 14

PIE_CHART_POSITION = "M2"
BAR_CHART_POSITION = "M20"


class DashboardCharts:
    """
    Dashboard grafiklerini oluşturan yardımcı sınıf.
    """

    @staticmethod
    def add_match_pie_chart(
        worksheet: Worksheet,
        summary: DashboardSummary,
    ) -> None:
        """
        Eşleşme durumunu gösteren pasta grafiğini oluşturur.
        """

        worksheet["N1"] = "Durum"
        worksheet["O1"] = "Kayıt"

        worksheet["N2"] = "Eşleşen"
        worksheet["O2"] = summary.matched_records

        worksheet["N3"] = "MKYS Eksik"
        worksheet["O3"] = summary.missing_in_mkys

        worksheet["N4"] = "TDMS Eksik"
        worksheet["O4"] = summary.missing_in_tdms

        labels = Reference(
            worksheet,
            min_col=14,
            min_row=2,
            max_row=4,
        )

        data = Reference(
            worksheet,
            min_col=15,
            min_row=1,
            max_row=4,
        )

        chart = PieChart()

        chart.title = "Giriş Hareketleri"

        chart.height = CHART_HEIGHT
        chart.width = CHART_WIDTH

        chart.add_data(
            data,
            titles_from_data=True,
        )

        chart.set_categories(
            labels,
        )

        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True

        worksheet.add_chart(
            chart,
            PIE_CHART_POSITION,
        )

    @staticmethod
    def add_consumption_bar_chart(
        dashboard_sheet: Worksheet,
        analytics_sheet: Worksheet,
    ) -> None:
        """
        Tüketim ve tutar farklarını gösteren sütun grafiği ekler.
        """

        labels = Reference(
            analytics_sheet,
            min_col=14,
            min_row=8,
            max_row=9,
        )

        data = Reference(
            analytics_sheet,
            min_col=15,
            min_row=7,
            max_row=9,
        )

        chart = BarChart()

        chart.type = "col"
        chart.style = 10

        chart.title = "MKYS - TDMS Fark Analizi"
        chart.y_axis.title = "Kayıt"
        chart.x_axis.title = "Kategori"

        chart.height = CHART_HEIGHT
        chart.width = CHART_WIDTH

        chart.add_data(
            data,
            titles_from_data=True,
        )

        chart.set_categories(
            labels,
        )

        # Kurumsal mavi renk
        chart.series[0].graphicalProperties = GraphicalProperties(
            solidFill="4472C4",
        )

        # Veri etiketleri
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True

        # Grid çizgilerini kaldır
        chart.y_axis.majorGridlines = None

        dashboard_sheet.add_chart(
            chart,
            BAR_CHART_POSITION,
        )
